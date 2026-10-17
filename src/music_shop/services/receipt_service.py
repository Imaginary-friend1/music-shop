from openpyxl import Workbook

from music_shop.models.receipt import Receipt

class ReceiptService:
    """Сервис для работы с чеками."""

    def export_to_excel(self, receipt: Receipt, filename: str) -> None:
        """Сохраняет чек в Excel-файл."""
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Чек"

        sheet["A1"] = "Чек"
        sheet["A2"] = "Номер чека"
        sheet["B2"] = receipt.receipt_number

        sheet["A3"] = "Способ оплаты"
        sheet["B3"] = receipt.payment_method

        sheet["A5"] = "Артикул"
        sheet["B5"] = "Наименование"
        sheet["C5"] = "Цена"

        row = 6

        for instrument in receipt.items:
            sheet.cell(row=row, column=1, value=instrument.article)
            sheet.cell(row=row, column=2, value=instrument.name)
            sheet.cell(row=row, column=3, value=instrument.price)
            row += 1

        sheet.cell(row=row, column=2, value="Итого")
        sheet.cell(row=row, column=3, value=receipt.get_total())

        workbook.save(filename)
